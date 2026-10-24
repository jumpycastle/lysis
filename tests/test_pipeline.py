import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).parent))

from canned import AGENT_SCRIPT, CANNED  # noqa: E402
from fake_client import FakeClient  # noqa: E402

from lysis_threatmodelagent.code_tools import RepoSandbox  # noqa: E402
from lysis_threatmodelagent.config import load_profile  # noqa: E402
from lysis_threatmodelagent.enforce import cvss31_score, valid_ref  # noqa: E402
from lysis_threatmodelagent.evals import score_case  # noqa: E402
from lysis_threatmodelagent.intake import load_bundle, parse_drawio  # noqa: E402
from lysis_threatmodelagent.llm import LLM  # noqa: E402
from lysis_threatmodelagent.pipeline import ReviewOptions, run_review  # noqa: E402
from lysis_threatmodelagent.schema import Final, tool_schema  # noqa: E402

CASE = ROOT / "evals/cases/labshare"


@pytest.fixture(scope="module")
def result(tmp_path_factory):
    profile = load_profile(ROOT / "profiles/default.yaml")
    bundle = load_bundle(CASE / "artifacts", context=(CASE / "context.md").read_text())
    client = FakeClient(CANNED, AGENT_SCRIPT)
    llm = LLM(client=client)
    out = tmp_path_factory.mktemp("runs")
    res = run_review(bundle, profile, llm, ReviewOptions(repo=str(CASE / "repo"), out_dir=str(out),
                                                         log=lambda m: None))
    res.client = client
    return res


def fixes(res, rule):
    return [e for e in res.enforced.log.entries if e["rule"] == rule]


# ------------------------------------------------------------------ unit

def test_cvss_known_vectors():
    assert cvss31_score("CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H") == 9.8
    assert cvss31_score("CVSS:3.1/AV:N/AC:L/PR:L/UI:N/S:U/C:H/I:N/A:N") == 6.5
    assert cvss31_score("CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:C/C:H/I:H/A:H") == 10.0
    assert cvss31_score("CVSS:3.1/AV:L/AC:L/PR:L/UI:N/S:U/C:N/I:N/A:N") == 0.0
    assert cvss31_score("garbage") is None


def test_framework_refs():
    assert valid_ref("OWASP A01:2021")
    assert valid_ref("ATT&CK T1078.004")
    assert valid_ref("NIST 800-53 AC-6(9)")
    assert not valid_ref("NIST 800-53 ZZ-9")
    assert not valid_ref("ISO 27001 A.9")


def test_drawio_parse():
    txt = parse_drawio((CASE / "artifacts/dfd.drawio").read_text())
    assert "CRO LIMS -> ALB  [webhook POST (result ready)]" in txt
    assert "Flask API  (inside: AWS VPC (private subnets))" in txt


def test_tool_schema_keeps_title_property():
    s = tool_schema(Final)
    finding = s["properties"]["findings"]["items"]
    assert "title" in finding["properties"]
    assert "$ref" not in json.dumps(s)


def test_sandbox_blocks_escape():
    box = RepoSandbox(CASE / "repo")
    with pytest.raises(PermissionError):
        box.read_file({"path": "../../../../etc/passwd"})
    assert "tenant_id" in box.grep({"pattern": "tenant_id"})


# ------------------------------------------------------------------ pipeline + enforcement

def test_pipeline_outputs(result):
    d = result.run_dir
    for f in ("findings.json", "report.md", "run.json", "system_model.json", "verification.json"):
        assert (d / f).exists(), f
    exp = json.loads((d / "findings.json").read_text())
    assert exp["epic"]["title"] == "Threat Model – LabShare"
    stories = [f for f in exp["findings"] if f["type"] == "Story"]
    assert len(stories) == 1
    s = stories[0]
    assert s["title"] == "Risk Assessment – LabShare"
    assert s["tags"] == ["Threat Model Agent", "Design Review"]
    assert s["tshirt_size"] == "XS" and s["severity"] == "Informational"


def test_sandbox_escape_in_agent_loop(result):
    msgs = result.client.messages.calls
    # find the tool_result for the passwd read
    blob = json.dumps([m.get("messages") for m in msgs], default=str)
    assert "path escapes repository" in blob


def test_fabricated_quote_removed_and_basis_downgraded(result):
    assert any("HMAC" in e["detail"] for e in fixes(result, "quote_not_found"))
    r2 = next(f for f in result.enforced.findings if f.id == "R2")
    assert r2.evidence_basis == "inference"


def test_real_quote_and_code_kept(result):
    r1 = next(f for f in result.enforced.findings if f.id == "R1")
    assert r1.evidence_quotes and r1.evidence_quotes[0].verified
    assert r1.code_evidence and r1.code_evidence[0].verified
    assert r1.evidence_basis == "code_verified"


def test_refs_standards_owner(result):
    r1 = next(f for f in result.enforced.findings if f.id == "R1")
    assert r1.framework_refs == ["OWASP A01:2021", "NIST 800-53 AC-3"]
    assert r1.suggested_owner == "System Owner"
    r2 = next(f for f in result.enforced.findings if f.id == "R2")
    assert r2.standards_cited == ["Third-Party Risk Standard"]


def test_missing_docs_severity_cap(result):
    r3 = next(f for f in result.enforced.findings if f.id == "R3")
    assert r3.severity == "Medium" and r3.confidence == "Medium"


def test_duplicate_merged(result):
    ids = [f.id for f in result.enforced.findings]
    assert "R4" not in ids and fixes(result, "duplicate")


def test_gates(result):
    gen = fixes(result, "gate_finding_generated")
    gated = {e["detail"].split()[1] for e in gen}
    # logging=No, rbac=Yes-without-evidence; ir_kill_switch + infra_hygiene/secrets/patch are
    # unanswered -> Unknown -> findings; strong_auth passes (Yes with evidence)
    assert gated == {"security_logging", "rbac", "ir_kill_switch",
                     "infra_hygiene", "secrets_management", "patch_management"}
    assert fixes(result, "gate_unanswered")


def test_decision_clamped_and_renumbered(result):
    d = result.enforced.decisions[0]
    assert d.decision == "Conditional Approval"
    story = result.enforced.export.findings[0]
    assert story.description.startswith("🚫 Conditional Approval — LabShare")
    assert "[R1]" not in story.description and "#1" in story.description
    assert story.priority == "High"


def test_actionability(result):
    exp = {f.title: f for f in result.enforced.export.findings if f.type == "Risk"}
    r1 = next(v for k, v in exp.items() if k.startswith("Export endpoint"))
    assert r1.actionability == "CreateTicket" and r1.cvss == 6.5
    assert "Threat Model" in r1.tags and "Multi-tenancy" in r1.tags


def test_eval_scoring(result):
    import yaml
    expected = yaml.safe_load((CASE / "expected.yaml").read_text())
    sc = score_case("labshare", expected, result)
    assert "export-cross-tenant" in sc.must_find_hit
    assert sc.decision_actual == "Conditional Approval"
    assert not sc.false_positives


def test_caching_markers(result):
    call = result.client.messages.calls[0]
    assert call["system"][-1].get("cache_control") == {"type": "ephemeral"}
    user_blocks = call["messages"][0]["content"]
    assert any(b.get("cache_control") for b in user_blocks[:-1])



def _text_pdf(path, pages):
    """Minimal PDF with extractable text on every page (no extra deps)."""
    objs, kids = [], []
    font_id = 3
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(None)  # pages placeholder
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i in range(pages):
        stream = f"BT /F1 12 Tf 72 700 Td (Pentest finding {i}: IDOR on /api/export) Tj ET".encode()
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        content_id = len(objs)
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
                    f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode())
        kids.append(len(objs))
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {pages} >>".encode()
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for n, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{off:010d} 00000 n \n".encode() for off in offsets)
    out += f"trailer << /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    path.write_bytes(bytes(out))


def test_pptx_and_pdf_modes(tmp_path):
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Deployment"
    slide.placeholders[1].text = "EKS in private subnets"
    slide.notes_slide.notes_text_frame.text = "API keys stored in SSM"
    prs.save(tmp_path / "deck.pptx")
    _text_pdf(tmp_path / "small.pdf", 3)
    _text_pdf(tmp_path / "big.pdf", 45)

    b = load_bundle(tmp_path)
    kinds = {a.name: a.kind for a in b.artifacts}
    assert kinds == {"deck.pptx": "pptx", "small.pdf": "pdf", "big.pdf": "pdf-text"}
    deck = next(a for a in b.artifacts if a.kind == "pptx")
    assert "EKS in private subnets" in deck.text and "Speaker notes: API keys stored in SSM" in deck.text
    big = next(a for a in b.artifacts if a.name == "big.pdf")
    assert "Pentest finding 44" in big.text
    assert not b.skipped

    assert all(a.kind != "pdf" for a in load_bundle(tmp_path, pdf_mode="text").artifacts)
    assert {a.kind for a in load_bundle(tmp_path, pdf_mode="native").artifacts} == {"pptx", "pdf"}


def test_text_reply_is_reprompted(result):
    draft_calls = [c for c in result.client.messages.calls
                   if c["tools"][0]["name"] == "submit_draft" and len(c["tools"]) == 1]
    assert len(draft_calls) == 2
    assert "Call the `submit_draft` tool now" in json.dumps(draft_calls[1]["messages"], default=str)


def test_prior_and_diagram(result):
    # prior assessment auto-detected from the filename, extracted, and overlaid
    assert result.prior is not None and len(result.prior.findings) == 2
    d = result.run_dir
    assert (d / "prior_findings.json").exists()
    assert (d / "threat-model.html").exists()
    htmltext = (d / "threat-model.html").read_text()
    assert "ACME-PT-LS-001" in htmltext          # prior finding embedded
    assert '"prior_overlap": "duplicate"' in (d / "run.json").read_text() or \
           any(f.prior_overlap == "duplicate" for f in result.enforced.findings)
    # the diagram data has findings and prior arrays
    import re, json as J
    m = re.search(r"const D=(\{.*?\});", htmltext, re.S)
    assert m
    data = J.loads(m.group(1))
    assert len(data["findings"]) == len(result.enforced.findings)
    assert len(data["prior"]) == 2
    assert data["nodes"], "diagram should have nodes from the system model"
    assert len(data["chains"]) >= 1, "attack chains should be rendered"
    assert any(st["k"] == "pt" for c in data["chains"] for st in c["steps"]), \
        "a chain should be able to include a prior (pentest) step"
    assert "chainpath" in htmltext and "Attack chains" in htmltext


def test_diagram_unit_no_prior():
    from lysis_threatmodelagent.diagram import build_diagram
    from lysis_threatmodelagent.schema import SystemModel
    import sys as _s; sys.path.insert(0, str(Path(__file__).parent))
    from canned import SYSTEM_MODEL
    m = SystemModel.model_validate(SYSTEM_MODEL)
    fs = [ff for ff in []]  # empty findings still renders
    html = build_diagram(m, fs, {"created_utc": "x", "tool_version": "0"}, None)
    assert "<svg" in html and "Threat Model" in html
